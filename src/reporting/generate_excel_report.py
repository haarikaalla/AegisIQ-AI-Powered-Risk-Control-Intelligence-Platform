"""
AegisIQ — Executive Excel Report Generator
=============================================
WHY THIS DELIVERABLE EXISTS (AWM M&T rationale):
  Control owners and testers live in Excel day-to-day — they want to filter,
  pivot, and sanity-check numbers themselves, not just view a static picture.
  This workbook mirrors that: raw data lives on "Data_*" sheets, and every
  summary number on "Exec Summary" / "Control Scorecard" is a live formula
  (SUMIFS/COUNTIFS/AVERAGEIFS/INDEX-MATCH) referencing those sheets — so if the
  underlying data sheet is refreshed with a new extract, the whole workbook
  recalculates rather than needing to be manually rebuilt.

Formula choices follow the platform's Excel constraints: SUMIFS/COUNTIFS/
AVERAGEIFS/INDEX/MATCH/IFERROR only (Excel-2007-era, no XLOOKUP/SORT/FILTER/
UNIQUE, which the recalculation engine used to validate this file cannot
evaluate).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

PROCESSED_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data", "processed")
OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "outputs")
FONT_NAME = "Calibri"

# ---- Palette (matches the dashboard's control-room theme for a consistent brand) ----
NAVY = "1F2A44"
NAVY_DARK = "141B2E"
SLATE = "5B6B85"
WHITE = "FFFFFF"
ZEBRA = "EDEFF5"

HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
BANNER_FILL = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=20, color=WHITE)
SUBTITLE_FONT = Font(name=FONT_NAME, size=10.5, color="B9C2D6")
LABEL_FONT = Font(name=FONT_NAME, bold=True, size=10.5, color=NAVY)
BODY_FONT = Font(name=FONT_NAME, size=10)
ZEBRA_FILL = PatternFill(start_color=ZEBRA, end_color=ZEBRA, fill_type="solid")

KPI_TILE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
KPI_LABEL_FONT = Font(name=FONT_NAME, size=9.5, color=SLATE)
KPI_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=24, color=NAVY)
KPI_VALUE_FONT_RED = Font(name=FONT_NAME, bold=True, size=24, color="C0392B")

THIN = Side(style="thin", color="D9D9E3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TILE_BORDER = Border(left=Side(style="thin", color="D9D9E3"), right=Side(style="thin", color="D9D9E3"),
                      top=Side(style="thin", color="D9D9E3"), bottom=Side(style="thin", color="D9D9E3"))

RAG_FILL = {
    "Red": PatternFill(start_color="F8CBCB", end_color="F8CBCB", fill_type="solid"),
    "Amber": PatternFill(start_color="FCE7B2", end_color="FCE7B2", fill_type="solid"),
    "Green": PatternFill(start_color="C9E7CE", end_color="C9E7CE", fill_type="solid"),
}
TAB_COLORS = {
    "Exec Summary": "1F2A44",
    "Control Scorecard": "2F5C8A",
    "KRI Trend": "3E7C8C",
    "Data_ControlResults": "9AA5B8",
    "Data_RiskScores": "9AA5B8",
    "Data_RemediationActions": "9AA5B8",
}


def banner(ws, title, subtitle, ncols=8):
    """A dark title banner across the top of a sheet, matching the dashboard theme."""
    for c in range(1, ncols + 1):
        ws.cell(row=1, column=c).fill = BANNER_FILL
        ws.cell(row=2, column=c).fill = BANNER_FILL
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 16
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center", indent=1)
    ws.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT
    ws.cell(row=2, column=1).alignment = Alignment(vertical="center", indent=1)


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def write_df(ws, df, start_row=1, start_col=1, header=True, zebra=True, pct_cols=None, wrap_cols=None):
    pct_cols = pct_cols or []
    wrap_cols = wrap_cols or []
    if header:
        for j, col in enumerate(df.columns):
            ws.cell(row=start_row, column=start_col + j, value=str(col))
        style_header_row(ws, start_row, len(df.columns))
        start_row += 1
    for i, (_, row) in enumerate(df.iterrows()):
        for j, val in enumerate(row):
            cell = ws.cell(row=start_row + i, column=start_col + j, value=(None if pd.isna(val) else val))
            cell.font = BODY_FONT
            cell.border = BORDER
            if zebra and i % 2 == 1:
                cell.fill = ZEBRA_FILL
            if df.columns[j] in pct_cols:
                cell.number_format = "0.0%"
            if df.columns[j] in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        if wrap_cols:
            ws.row_dimensions[start_row + i].height = 50
    return start_row + len(df)  # next free row


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_table(ws, ref, name, style="TableStyleMedium9"):
    """Adds a named Excel Table (banded rows + built-in autofilter) for a polished, sortable feel."""
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showFirstColumn=False)
    ws.add_table(tbl)


def main():
    results = pd.read_csv(os.path.join(PROCESSED_DIR, "fact_control_test_results.csv"))
    risk = pd.read_csv(os.path.join(PROCESSED_DIR, "fact_risk_scores.csv"))
    remediation = pd.read_csv(os.path.join(PROCESSED_DIR, "remediation_actions.csv"))
    dq_log = pd.read_csv(os.path.join(PROCESSED_DIR, "data_quality_log.csv"))
    catalog = pd.read_csv(os.path.join(PROCESSED_DIR, "control_catalog_clean.csv"))

    wb = Workbook()

    # =========================================================================
    # SHEET: Data_ControlResults (raw, full history — feeds formulas elsewhere)
    # =========================================================================
    ws_data = wb.active
    ws_data.title = "Data_ControlResults"
    ws_data.sheet_view.showGridLines = False
    cols = ["control_id", "control_name", "process_name", "period", "population",
            "pass_count", "fail_count", "dq_exception_count", "pass_rate",
            "kri_threshold", "rag_status"]
    end_row = write_df(ws_data, results[cols], zebra=False, pct_cols=["pass_rate", "kri_threshold"])
    autosize(ws_data, [10, 26, 30, 10, 12, 10, 10, 14, 10, 12, 10])
    ws_data.freeze_panes = "A2"
    add_table(ws_data, f"A1:{get_column_letter(len(cols))}{end_row-1}", "TblControlResults")
    n_results = len(results)

    # =========================================================================
    # SHEET: Data_RiskScores
    # =========================================================================
    ws_risk = wb.create_sheet("Data_RiskScores")
    ws_risk.sheet_view.showGridLines = False
    risk_cols = ["control_id", "control_name", "process_name", "risk_category", "period",
                 "latest_pass_rate", "kri_threshold", "inherent_risk_1to5",
                 "control_effectiveness_score", "control_effectiveness_rating",
                 "residual_risk_score", "residual_risk_rating", "open_exceptions", "rag_status_latest"]
    end_row = write_df(ws_risk, risk[risk_cols], zebra=False, pct_cols=["latest_pass_rate", "kri_threshold"])
    autosize(ws_risk, [10, 26, 30, 14, 10, 12, 10, 10, 14, 18, 12, 14, 10, 10])
    ws_risk.freeze_panes = "A2"
    add_table(ws_risk, f"A1:{get_column_letter(len(risk_cols))}{end_row-1}", "TblRiskScores")
    n_risk = len(risk)

    # =========================================================================
    # SHEET: Data_RemediationActions
    # =========================================================================
    ws_rem = wb.create_sheet("Data_RemediationActions")
    ws_rem.sheet_view.showGridLines = False
    rem_cols = ["action_id", "exception_id", "control_id", "owner", "reason_code", "detail",
                "identified_date", "target_closure_date", "days_to_target", "status_vs_sla",
                "remediation_priority"]
    end_row = write_df(ws_rem, remediation[rem_cols], zebra=False)
    autosize(ws_rem, [14, 12, 10, 24, 20, 45, 14, 16, 12, 12, 14])
    ws_rem.freeze_panes = "A2"
    add_table(ws_rem, f"A1:{get_column_letter(len(rem_cols))}{end_row-1}", "TblRemediation")
    n_rem = len(remediation)

    # =========================================================================
    # SHEET: Control Scorecard (formula-driven, one row per control, latest period)
    # =========================================================================
    ws_sc = wb.create_sheet("Control Scorecard", 0)
    ws_sc.sheet_view.showGridLines = False
    ncols_sc = 11
    banner(ws_sc, "AegisIQ — Control Scorecard",
           "Latest tested period per control, pulled live from Data_ControlResults / Data_RiskScores",
           ncols=ncols_sc)

    header = ["Control ID", "Control Name", "Process", "Latest Period", "Population",
              "Pass Rate", "KRI Threshold", "RAG Status", "Control Effectiveness",
              "Residual Risk", "Open Exceptions"]
    row0 = 4
    for j, h in enumerate(header):
        ws_sc.cell(row=row0, column=j + 1, value=h)
    style_header_row(ws_sc, row0, len(header))

    control_ids = catalog["control_id"].tolist()
    r_ctrl_data = n_results + 1  # +1 for header row on Data_ControlResults
    r_risk_data = n_risk + 1

    for i, cid in enumerate(control_ids):
        r = row0 + 1 + i
        ws_sc.cell(row=r, column=1, value=cid).font = Font(name=FONT_NAME, bold=True, size=10, color=NAVY)
        ws_sc.cell(row=r, column=2,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$B$2:$B${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=3,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$C$2:$C${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=4,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$E$2:$E${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=5,
                   value=f'=IFERROR(SUMIFS(Data_ControlResults!$E$2:$E${r_ctrl_data},Data_ControlResults!$A$2:$A${r_ctrl_data},$A{r},Data_ControlResults!$D$2:$D${r_ctrl_data},$D{r}),"")')
        ws_sc.cell(row=r, column=6,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$F$2:$F${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=6).number_format = "0.0%"
        ws_sc.cell(row=r, column=7,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$G$2:$G${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=7).number_format = "0.0%"
        ws_sc.cell(row=r, column=8,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$N$2:$N${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=9,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$I$2:$I${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=10,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$L$2:$L${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        ws_sc.cell(row=r, column=11,
                   value=f'=IFERROR(INDEX(Data_RiskScores!$M$2:$M${r_risk_data},MATCH($A{r},Data_RiskScores!$A$2:$A${r_risk_data},0)),"")')
        for c in range(1, len(header) + 1):
            ws_sc.cell(row=r, column=c).border = BORDER
            if c not in (1, 6, 7):
                ws_sc.cell(row=r, column=c).font = BODY_FONT
            if i % 2 == 1:
                ws_sc.cell(row=r, column=c).fill = ZEBRA_FILL

    last_data_row = row0 + len(control_ids)
    ws_sc.freeze_panes = f"A{row0+1}"
    ws_sc.auto_filter.ref = f"A{row0}:K{last_data_row}"

    ws_sc.conditional_formatting.add(
        f"H{row0+1}:H{last_data_row}",
        CellIsRule(operator="equal", formula=['"Red"'], fill=RAG_FILL["Red"]))
    ws_sc.conditional_formatting.add(
        f"H{row0+1}:H{last_data_row}",
        CellIsRule(operator="equal", formula=['"Amber"'], fill=RAG_FILL["Amber"]))
    ws_sc.conditional_formatting.add(
        f"H{row0+1}:H{last_data_row}",
        CellIsRule(operator="equal", formula=['"Green"'], fill=RAG_FILL["Green"]))
    ws_sc.conditional_formatting.add(
        f"J{row0+1}:J{last_data_row}",
        CellIsRule(operator="equal", formula=['"High"'], fill=RAG_FILL["Amber"]))
    ws_sc.conditional_formatting.add(
        f"J{row0+1}:J{last_data_row}",
        CellIsRule(operator="equal", formula=['"Critical"'], fill=RAG_FILL["Red"]))
    ws_sc.conditional_formatting.add(
        f"J{row0+1}:J{last_data_row}",
        CellIsRule(operator="equal", formula=['"Low"'], fill=RAG_FILL["Green"]))

    autosize(ws_sc, [10, 26, 30, 12, 10, 10, 12, 10, 16, 12, 14])

    # =========================================================================
    # SHEET: Exec Summary (formula-driven KPI tiles)
    # =========================================================================
    ws_es = wb.create_sheet("Exec Summary", 0)
    ws_es.sheet_view.showGridLines = False
    ncols_es = 8
    banner(ws_es, "AegisIQ — Executive Summary",
           "AWM Monitoring & Testing Program  |  Synthetic data, demonstration purposes",
           ncols=ncols_es)

    # ---- KPI tile row (2x4 grid of merged-cell tiles with borders, like a BI card strip) ----
    kpi_defs = [
        ("CONTROLS IN SCOPE", f"=COUNTA('Control Scorecard'!A5:A{last_data_row})", False),
        ("RATED RED (latest)", f"=COUNTIF('Control Scorecard'!H5:H{last_data_row},\"Red\")", True),
        ("RATED AMBER (latest)", f"=COUNTIF('Control Scorecard'!H5:H{last_data_row},\"Amber\")", False),
        ("RATED GREEN (latest)", f"=COUNTIF('Control Scorecard'!H5:H{last_data_row},\"Green\")", False),
        ("HIGH/CRITICAL RESIDUAL RISK",
         f"=COUNTIF('Control Scorecard'!J5:J{last_data_row},\"High\")+COUNTIF('Control Scorecard'!J5:J{last_data_row},\"Critical\")", True),
        ("OPEN REMEDIATION ACTIONS", f"=COUNTA(Data_RemediationActions!A2:A{n_rem+1})", False),
        ("OVERDUE VS. SLA", f'=COUNTIF(Data_RemediationActions!J2:J{n_rem+1},"Overdue")', True),
        ("AVG PASS RATE (latest)", f"=AVERAGE('Control Scorecard'!F5:F{last_data_row})", False),
    ]
    tile_row0 = 4
    tile_h = 4  # rows tall per tile
    col_per_tile = 2  # columns wide per tile
    for idx, (label, formula, is_red) in enumerate(kpi_defs):
        col_start = 1 + (idx % 4) * col_per_tile
        row_start = tile_row0 + (idx // 4) * (tile_h + 1)
        row_end = row_start + tile_h - 1
        col_end = col_start + col_per_tile - 1
        for rr in range(row_start, row_end + 1):
            for cc in range(col_start, col_end + 1):
                ws_es.cell(row=rr, column=cc).fill = KPI_TILE_FILL
                ws_es.cell(row=rr, column=cc).border = TILE_BORDER
        ws_es.merge_cells(start_row=row_start, start_column=col_start, end_row=row_start, end_column=col_end)
        lbl_cell = ws_es.cell(row=row_start, column=col_start, value=label)
        lbl_cell.font = KPI_LABEL_FONT
        lbl_cell.alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)

        ws_es.merge_cells(start_row=row_start + 1, start_column=col_start, end_row=row_end, end_column=col_end)
        val_cell = ws_es.cell(row=row_start + 1, column=col_start, value=formula)
        val_cell.font = KPI_VALUE_FONT_RED if is_red else KPI_VALUE_FONT
        val_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if label == "AVG PASS RATE (latest)":
            val_cell.number_format = "0.0%"

    dq_start_row = tile_row0 + 2 * (tile_h + 1) + 1
    ws_es.cell(row=dq_start_row, column=1, value="Data Quality Log — Ingestion & Cleaning Layer").font = LABEL_FONT
    write_df(ws_es, dq_log, start_row=dq_start_row + 1, zebra=True, wrap_cols=["action_taken", "issue"])
    autosize(ws_es, [16, 24, 12, 24, 16, 16, 16, 16])

    # =========================================================================
    # SHEET: KRI Trend (control x period pivot, formula-driven, with chart)
    # =========================================================================
    ws_tr = wb.create_sheet("KRI Trend")
    ws_tr.sheet_view.showGridLines = False
    periods = sorted(results["period"].unique())
    banner(ws_tr, "AegisIQ — KRI Pass-Rate Trend", "Monthly pass rate by control, live from Data_ControlResults",
           ncols=len(periods) + 1)
    hdr_row = 4
    ws_tr.cell(row=hdr_row, column=1, value="Control ID")
    for j, p in enumerate(periods):
        ws_tr.cell(row=hdr_row, column=2 + j, value=p)
    style_header_row(ws_tr, hdr_row, len(periods) + 1)
    for i, cid in enumerate(control_ids):
        r = hdr_row + 1 + i
        ws_tr.cell(row=r, column=1, value=cid).font = Font(name=FONT_NAME, bold=True, size=10, color=NAVY)
        for j, p in enumerate(periods):
            period_ref = ws_tr.cell(row=hdr_row, column=2 + j).coordinate.replace(str(hdr_row), f"${hdr_row}")
            f = (f'=IFERROR(AVERAGEIFS(Data_ControlResults!$I$2:$I${r_ctrl_data},'
                 f'Data_ControlResults!$A$2:$A${r_ctrl_data},$A{r},'
                 f'Data_ControlResults!$D$2:$D${r_ctrl_data},{period_ref}),"")')
            cell = ws_tr.cell(row=r, column=2 + j, value=f)
            cell.number_format = "0.0%"
            cell.font = BODY_FONT
            cell.border = BORDER
            if i % 2 == 1:
                cell.fill = ZEBRA_FILL
    ws_tr.freeze_panes = ws_tr.cell(row=hdr_row + 1, column=2).coordinate
    autosize(ws_tr, [10] + [10] * len(periods))

    chart = LineChart()
    chart.title = "Pass Rate Trend by Control"
    chart.y_axis.title = "Pass Rate"
    chart.x_axis.title = "Period"
    chart.height, chart.width = 11, 28
    chart.style = 2
    data = Reference(ws_tr, min_col=2, max_col=1 + len(periods), min_row=hdr_row, max_row=hdr_row + len(control_ids))
    chart.add_data(data, titles_from_data=False, from_rows=True)
    chart.set_categories(Reference(ws_tr, min_col=2, max_col=1 + len(periods), min_row=hdr_row, max_row=hdr_row))
    ws_tr.add_chart(chart, f"A{hdr_row + 2 + len(control_ids)}")

    # =========================================================================
    # Workbook-level polish: tab order, tab colors, active sheet
    # =========================================================================
    wb._sheets = [wb["Exec Summary"], wb["Control Scorecard"], ws_tr,
                  wb["Data_ControlResults"], wb["Data_RiskScores"], wb["Data_RemediationActions"]]
    for name, color in TAB_COLORS.items():
        wb[name].sheet_properties.tabColor = color
    wb.active = 0

    os.makedirs(OUTPUTS_DIR, exist_ok=True)
    out_path = os.path.join(OUTPUTS_DIR, "AegisIQ_Executive_Report.xlsx")
    wb.save(out_path)
    print(f"Workbook written to {out_path}")


if __name__ == "__main__":
    main()
